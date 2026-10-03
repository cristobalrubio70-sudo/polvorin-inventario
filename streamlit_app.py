import streamlit as st
import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import date
import uuid
import warnings
import io
warnings.filterwarnings("ignore")

# ── Configuracion ─────────────────────────────────────────────────────────────
st.set_page_config(
    page_title="Analizador Polvorin Minero",
    page_icon="🧨",
    layout="wide"
)

EXPLOSIVOS_CONFIG = {
    "ANFO":               {"tipo":"Secundario","unidad":"saco",  "peso_kg":25.0,"vida_dias":365,"stock_min":5, "almacen":"Container secundarios"},
    "Emulsion":           {"tipo":"Secundario","unidad":"caja",  "peso_kg":20.0,"vida_dias":180,"stock_min":3, "almacen":"Container secundarios"},
    "Mecha de seguridad": {"tipo":"Secundario","unidad":"rollo", "peso_kg":5.0, "vida_dias":730,"stock_min":2, "almacen":"Container secundarios"},
    "Fulminante N8":      {"tipo":"Primario",  "unidad":"caja",  "peso_kg":0.5, "vida_dias":365,"stock_min":2, "almacen":"Boveda enterrada (primarios)"}
}

NOMBRE_MAP = {
    "Emulsion":"Emulsion","Emulsión":"Emulsion",
    "Fulminante N°8":"Fulminante N8","Fulminante N8":"Fulminante N8",
    "ANFO":"ANFO","Mecha de seguridad":"Mecha de seguridad",
}

COLS_MIN = [
    "id_lote","explosivo","tipo","cantidad","unidad",
    "fecha_ingreso","fecha_vencimiento","proveedor",
    "n_guia","operador_ingreso","almacen"
]

# ── Funciones ─────────────────────────────────────────────────────────────────
def cargar_archivo(archivo):
    errores = []
    mov = pd.DataFrame()
    try:
        nombre = archivo.name.lower()
        if nombre.endswith(".xlsx") or nombre.endswith(".xls"):
            xl = pd.ExcelFile(archivo)
            hojas = xl.sheet_names
            hoja_inv = "Inventario" if "Inventario" in hojas else hojas[0]
            if hoja_inv != "Inventario":
                errores.append("Hoja Inventario no encontrada. Se uso: " + hoja_inv)
            df = pd.read_excel(archivo, sheet_name=hoja_inv, header=3)
            if "Movimientos" in hojas:
                mov = pd.read_excel(archivo, sheet_name="Movimientos", header=3)
        elif nombre.endswith(".csv"):
            df = pd.read_csv(archivo)
        else:
            return None, None, ["Formato no soportado. Use .xlsx o .csv"]

        df.columns = [str(c).strip().lower().replace(" ","_") for c in df.columns]
        df = df.dropna(how="all")
        if "explosivo" in df.columns:
            df["explosivo"] = df["explosivo"].apply(
                lambda x: NOMBRE_MAP.get(str(x).strip(), str(x).strip()))
        falt = [c for c in COLS_MIN if c not in df.columns]
        if falt:
            errores.append("Columnas faltantes: " + ", ".join(falt))
        return df, mov, errores
    except Exception as e:
        return None, None, ["Error al leer archivo: " + str(e)]


def clasificar_estado(row):
    exp = str(row.get("explosivo","")).strip()
    cfg = EXPLOSIVOS_CONFIG.get(exp, {})
    hoy = date.today()
    estados = []
    try:
        fv   = pd.to_datetime(str(row.get("fecha_vencimiento",""))).date()
        dias = (fv - hoy).days
        if dias < 0:
            estados.append("Vencido")
        elif dias <= 30:
            estados.append("Por vencer")
    except Exception:
        estados.append("Fecha invalida")
    try:
        cant = int(float(str(row.get("cantidad", 0))))
        smin = cfg.get("stock_min", 0)
        if smin and cant < smin:
            estados.append("Stock bajo")
    except Exception:
        pass
    alm_ok  = cfg.get("almacen","")
    alm_act = str(row.get("almacen","")).strip()
    if alm_ok and alm_act and alm_act != alm_ok:
        estados.append("Almacen incorrecto")
    if exp and exp not in EXPLOSIVOS_CONFIG:
        estados.append("Explosivo desconocido")
    try:
        cant = int(float(str(row.get("cantidad", 0))))
        pku  = cfg.get("peso_kg", 0)
        pr   = float(str(row.get("peso_kg_total", 0)))
        pe   = round(cant * pku, 2)
        if pku and pe > 0 and abs(pr - pe) > pe * 0.15:
            estados.append("Peso inconsistente")
    except Exception:
        pass
    return "Vigente" if not estados else " | ".join(estados)


def evaluar(df):
    hoy = date.today()
    nuevas = []
    for _, row in df.iterrows():
        exp  = str(row.get("explosivo","")).strip()
        cfg  = EXPLOSIVOS_CONFIG.get(exp, {})
        lote = str(row.get("id_lote","?"))
        try:
            fv   = pd.to_datetime(str(row.get("fecha_vencimiento",""))).date()
            dias = (fv - hoy).days
            if dias < 0:
                nuevas.append({"id_disc":"D-"+uuid.uuid4().hex[:5].upper(),
                    "id_lote":lote,"explosivo":exp,"tipo":"Vencimiento superado",
                    "esperado":"Fecha futura","encontrado":str(fv),
                    "dias":abs(dias),"gravedad":"CRITICA","estado":"Pendiente",
                    "accion":"Destruccion autorizada - Ley de Armas Cap. VIII"})
            elif dias <= 30:
                nuevas.append({"id_disc":"D-"+uuid.uuid4().hex[:5].upper(),
                    "id_lote":lote,"explosivo":exp,"tipo":"Vencimiento proximo (<30 dias)",
                    "esperado":">30 dias","encontrado":str(dias)+" dias restantes",
                    "dias":dias,"gravedad":"ALTA","estado":"En proceso",
                    "accion":"Notificar a Autoridad Fiscalizadora - DGMN"})
        except Exception:
            nuevas.append({"id_disc":"D-"+uuid.uuid4().hex[:5].upper(),
                "id_lote":lote,"explosivo":exp,"tipo":"Fecha vencimiento invalida",
                "esperado":"AAAA-MM-DD","encontrado":str(row.get("fecha_vencimiento","")),
                "dias":0,"gravedad":"MEDIA","estado":"Pendiente",
                "accion":"Corregir registro manualmente"})
        try:
            cant = int(float(str(row.get("cantidad", 0))))
            smin = cfg.get("stock_min", 0)
            if smin and cant < smin:
                nuevas.append({"id_disc":"D-"+uuid.uuid4().hex[:5].upper(),
                    "id_lote":lote,"explosivo":exp,"tipo":"Stock bajo minimo operacional",
                    "esperado":">="+str(smin)+" "+cfg.get("unidad",""),
                    "encontrado":str(cant)+" "+str(row.get("unidad","")),
                    "dias":0,"gravedad":"MEDIA","estado":"Pendiente",
                    "accion":"Solicitar reabastecimiento al proveedor"})
        except Exception:
            pass
        alm_ok  = cfg.get("almacen","")
        alm_act = str(row.get("almacen","")).strip()
        if alm_ok and alm_act and alm_act != alm_ok:
            nuevas.append({"id_disc":"D-"+uuid.uuid4().hex[:5].upper(),
                "id_lote":lote,"explosivo":exp,"tipo":"Almacen incorrecto",
                "esperado":alm_ok,"encontrado":alm_act,
                "dias":0,"gravedad":"CRITICA","estado":"Pendiente",
                "accion":"Reubicar inmediatamente - DS 132 / Ley 17.798"})
        if exp and exp not in EXPLOSIVOS_CONFIG:
            nuevas.append({"id_disc":"D-"+uuid.uuid4().hex[:5].upper(),
                "id_lote":lote,"explosivo":exp,"tipo":"Explosivo no reconocido",
                "esperado":str(list(EXPLOSIVOS_CONFIG.keys())),"encontrado":exp,
                "dias":0,"gravedad":"MEDIA","estado":"Pendiente",
                "accion":"Verificar nomenclatura del explosivo"})
        try:
            cant = int(float(str(row.get("cantidad", 0))))
            pku  = cfg.get("peso_kg", 0)
            pr   = float(str(row.get("peso_kg_total", 0)))
            pe   = round(cant * pku, 2)
            if pku and pe > 0 and abs(pr - pe) > pe * 0.15:
                nuevas.append({"id_disc":"D-"+uuid.uuid4().hex[:5].upper(),
                    "id_lote":lote,"explosivo":exp,"tipo":"Peso inconsistente con cantidad",
                    "esperado":str(pe)+" kg","encontrado":str(pr)+" kg",
                    "dias":0,"gravedad":"BAJA","estado":"Revisar",
                    "accion":"Verificar pesaje del lote"})
        except Exception:
            pass
    cols = ["id_disc","id_lote","explosivo","tipo","esperado",
            "encontrado","dias","gravedad","estado","accion"]
    return pd.DataFrame(nuevas) if nuevas else pd.DataFrame(columns=cols)


def generar_excel(df, disc, mov):
    output = io.BytesIO()
    wb   = openpyxl.Workbook()
    HF   = PatternFill("solid", fgColor="1F3864")
    FF   = Font(bold=True, color="FFFFFF", size=10)
    BD   = Border(left=Side(style="thin"),right=Side(style="thin"),
                  top=Side(style="thin"), bottom=Side(style="thin"))
    FILL = {
        "CRITICA":   PatternFill("solid",fgColor="FFCCCC"),
        "ALTA":      PatternFill("solid",fgColor="FFEB9C"),
        "MEDIA":     PatternFill("solid",fgColor="FFF2CC"),
        "BAJA":      PatternFill("solid",fgColor="E2EFDA"),
        "VIGENTE":   PatternFill("solid",fgColor="C6EFCE"),
        "VENCIDO":   PatternFill("solid",fgColor="FFCCCC"),
        "POR VENCER":PatternFill("solid",fgColor="FFEB9C"),
        "STOCK":     PatternFill("solid",fgColor="FFF2CC"),
        "ALMACEN":   PatternFill("solid",fgColor="F4CCFF"),
    }
    def escribir(ws, data, titulo, sub=""):
        ws.append([titulo])
        ws.cell(1,1).font = Font(bold=True,size=13,color="1F3864")
        ws.append([sub or "Analisis: "+str(date.today())])
        ws.cell(2,1).font = Font(italic=True,size=10,color="595959")
        ws.append([])
        ws.append(list(data.columns))
        for ci,_ in enumerate(data.columns,1):
            c=ws.cell(4,ci); c.fill=HF; c.font=FF; c.border=BD
            c.alignment=Alignment(horizontal="center",wrap_text=True)
        for ri,row in enumerate(data.itertuples(index=False),5):
            for ci,val in enumerate(row,1):
                c=ws.cell(ri,ci,value=str(val) if val is not None else "")
                c.border=BD; c.alignment=Alignment(wrap_text=True,vertical="center")
                vs=str(val).upper()
                for k,f in FILL.items():
                    if k in vs: c.fill=f; break
        for col in ws.columns:
            ml=max((len(str(c.value)) if c.value else 0) for c in col)
            ws.column_dimensions[get_column_letter(col[0].column)].width=min(ml+4,40)

    ws1=wb.active; ws1.title="Inventario clasificado"
    escribir(ws1,df,"INVENTARIO ANALIZADO Y CLASIFICADO",
             "Lotes: "+str(len(df))+" | Analisis: "+str(date.today()))
    ws2=wb.create_sheet("Discrepancias")
    nc=len(disc[disc.gravedad=="CRITICA"]) if len(disc) else 0
    na=len(disc[disc.gravedad=="ALTA"])    if len(disc) else 0
    escribir(ws2,disc,"DISCREPANCIAS DETECTADAS",
             "Total: "+str(len(disc))+" | Criticas: "+str(nc)+" | Altas: "+str(na))
    if not mov.empty:
        ws3=wb.create_sheet("Movimientos")
        escribir(ws3,mov,"MOVIMIENTOS REGISTRADOS","Total: "+str(len(mov)))
    ws4=wb.create_sheet("Resumen ejecutivo")
    def ce(kw):
        if "estado" not in df.columns: return "N/A"
        return len(df[df["estado"].str.upper().str.contains(kw,na=False)])
    nc2=len(disc[disc.gravedad=="CRITICA"]) if len(disc) else 0
    na2=len(disc[disc.gravedad=="ALTA"])    if len(disc) else 0
    nm2=len(disc[disc.gravedad=="MEDIA"])   if len(disc) else 0
    nb2=len(disc[disc.gravedad=="BAJA"])    if len(disc) else 0
    for fila in [
        ["RESUMEN EJECUTIVO - REPORTE POLVORIN",""],
        ["Fecha de analisis",str(date.today())],["",""],
        ["INVENTARIO",""],
        ["Total lotes",len(df)],["Vigentes",ce("VIGENTE")],
        ["Vencidos",ce("VENCIDO")],["Por vencer (<30d)",ce("POR VENCER")],
        ["Stock bajo",ce("STOCK")],["Almacen incorrecto",ce("ALMACEN")],
        ["",""],["DISCREPANCIAS",""],
        ["Total",len(disc)],["Criticas",nc2],["Altas",na2],
        ["Medias",nm2],["Bajas",nb2],["",""],
        ["NORMATIVA",""],
        ["DS 132 SERNAGEOMIN","Reglamento de Seguridad Minera"],
        ["Ley 17.798","Control de Armas y Explosivos"],
        ["DGMN 003048/2020","Almacenamiento productos Ley 17.798"],
    ]:
        ws4.append(fila)
    ws4["A1"].font=Font(bold=True,size=13,color="1F3864")
    ws4.column_dimensions["A"].width=38
    ws4.column_dimensions["B"].width=50
    wb.save(output)
    output.seek(0)
    return output


# ── Interfaz Streamlit ────────────────────────────────────────────────────────
st.title("🧨 Analizador de Inventario — Polvorín Minero")
st.markdown("**Memoria de Título · Cristóbal Rubio · UNAB 2026**")
st.markdown(
    "> Normativa base: DS N° 132 SERNAGEOMIN · Ley N° 17.798 · Resolución DGMN N° 003048/2020"
)
st.divider()

tab1, tab2, tab3, tab4 = st.tabs(
    ["📤 Cargar y analizar", "📦 Inventario", "⚠️ Discrepancias", "📋 Movimientos"]
)

with tab1:
    st.subheader("Sube el archivo de inventario crudo")
    st.info("El archivo NO necesita columna de estado — el sistema la genera automáticamente.")
    archivo = st.file_uploader(
        "Selecciona el archivo .xlsx generado por el Notebook B (o un .csv)",
        type=["xlsx","xls","csv"]
    )

    if archivo:
        with st.spinner("Cargando y analizando..."):
            df, mov, errores = cargar_archivo(archivo)

        if df is None:
            for e in errores:
                st.error(e)
        else:
            if errores:
                for e in errores:
                    st.warning(e)

            df["estado"] = df.apply(clasificar_estado, axis=1)
            disc = evaluar(df)

            st.session_state["inv"]  = df
            st.session_state["disc"] = disc
            st.session_state["mov"]  = mov

            # Métricas en columnas
            def ce(kw):
                return len(df[df["estado"].str.upper().str.contains(kw,na=False)])

            nd = len(disc)
            nc = len(disc[disc.gravedad=="CRITICA"]) if nd else 0
            na = len(disc[disc.gravedad=="ALTA"])    if nd else 0

            col1,col2,col3,col4,col5 = st.columns(5)
            col1.metric("Total lotes",     len(df))
            col2.metric("✅ Vigentes",      ce("VIGENTE"))
            col3.metric("❌ Vencidos",      ce("VENCIDO"),  delta_color="inverse")
            col4.metric("⚠️ Por vencer",    ce("POR VENCER"), delta_color="inverse")
            col5.metric("🚨 Discrepancias", nd, delta_color="inverse")

            st.divider()

            if nc > 0:
                st.error(str(nc) + " discrepancia(s) CRÍTICA(S) — acción inmediata requerida.")
            elif na > 0:
                st.warning(str(na) + " discrepancia(s) de alta prioridad.")
            else:
                st.success("Sin discrepancias críticas.")

            # Boton descarga Excel
            excel_bytes = generar_excel(df, disc, mov)
            st.download_button(
                label="📥 Descargar reporte Excel completo",
                data=excel_bytes,
                file_name="reporte_polvorin_" + str(date.today()) + ".xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )

with tab2:
    st.subheader("Inventario clasificado")
    if "inv" in st.session_state and not st.session_state["inv"].empty:
        df_show = st.session_state["inv"]

        # Filtro rapido por estado
        estados_unicos = ["Todos"] + sorted(df_show["estado"].unique().tolist())
        filtro = st.selectbox("Filtrar por estado:", estados_unicos)
        if filtro != "Todos":
            df_show = df_show[df_show["estado"] == filtro]

        st.dataframe(df_show, use_container_width=True, hide_index=True)
        st.caption("Total mostrado: " + str(len(df_show)) + " lotes")
    else:
        st.info("Carga un archivo en la pestaña anterior para ver el inventario.")

with tab3:
    st.subheader("Discrepancias detectadas automáticamente")
    st.markdown("**CRITICA** = acción inmediata · **ALTA** = urgente · **MEDIA** = seguimiento · **BAJA** = revisar")
    if "disc" in st.session_state and not st.session_state["disc"].empty:
        disc_show = st.session_state["disc"]

        # Filtro por gravedad
        gravedades = ["Todas"] + sorted(disc_show["gravedad"].unique().tolist())
        filtro_g = st.selectbox("Filtrar por gravedad:", gravedades)
        if filtro_g != "Todas":
            disc_show = disc_show[disc_show["gravedad"] == filtro_g]

        # Color por gravedad
        def colorear(val):
            colores = {
                "CRITICA": "background-color: #FFCCCC",
                "ALTA":    "background-color: #FFEB9C",
                "MEDIA":   "background-color: #FFF2CC",
                "BAJA":    "background-color: #E2EFDA",
            }
            return colores.get(str(val).upper(), "")

        cols_d = ["id_lote","explosivo","tipo","gravedad","esperado","encontrado","accion"]
        disc_filtrada = disc_show[cols_d] if all(c in disc_show.columns for c in cols_d) else disc_show
        st.dataframe(
            disc_filtrada.style.applymap(colorear, subset=["gravedad"]),
            use_container_width=True,
            hide_index=True
        )
        st.caption("Total mostrado: " + str(len(disc_filtrada)) + " discrepancias")
    else:
        st.info("Carga un archivo en la pestaña anterior para ver las discrepancias.")

with tab4:
    st.subheader("Movimientos registrados")
    if "mov" in st.session_state and not st.session_state["mov"].empty:
        st.dataframe(st.session_state["mov"], use_container_width=True, hide_index=True)
        st.caption("Total: " + str(len(st.session_state["mov"])) + " movimientos")
    else:
        st.info("No hay movimientos en el archivo cargado (o aun no se ha cargado un archivo).")

st.divider()
st.caption(
    "Sistema de Inventario Polvorin Minero · Memoria de Titulo UNAB 2026 · "
    "Normativa: DS 132 | Ley 17.798 | DGMN 003048/2020"
)
